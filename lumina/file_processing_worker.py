"""Bounded local parsers. Executed as a script, never imports application actions."""
import base64
import csv
import io
import json
import math
import re
import stat
import sys
import zipfile
from pathlib import PurePosixPath
from xml.etree import ElementTree


# Source files are decoded as data only; no importing or execution is allowed.
TEXT_FORMATS = {'.txt', '.md', '.rst', '.log', '.csv', '.tsv', '.json', '.xml',
                '.py', '.js', '.ts', '.jsx', '.tsx', '.html', '.css', '.java',
                '.c', '.h', '.cpp', '.cs', '.go', '.rs', '.rb', '.php', '.swift',
                '.kt', '.sh', '.bash', '.ps1', '.lua', '.r', '.m', '.sql',
                '.yaml', '.yml', '.toml'}


class Rejected(ValueError):
    pass


def strict_json(text):
    def constant(value):
        raise Rejected('Nonfinite JSON numbers are unsupported')

    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise Rejected('Duplicate JSON object keys are unsupported')
            result[key] = value
        return result

    def number(value):
        parsed = float(value)
        if not math.isfinite(parsed):
            raise Rejected('Nonfinite JSON numbers are unsupported')
        return parsed
    return json.loads(text, parse_constant=constant, object_pairs_hook=pairs, parse_float=number)


def transform(data, suffix, action, limits, options, output_suffix):
    """Deterministic bounded operations, returning bytes to the parent for publication."""
    allowed = {'stats': set(), 'validate': set(), 'filter': {'column', 'condition', 'value'},
               'sort': {'column', 'numeric', 'ascending'}, 'format': {'indent', 'sort_keys'},
               'resize': {'width', 'height', 'quality'}, 'convert': {'quality'}}
    options = {} if options is None else options
    if not isinstance(options, dict) or options.keys() - allowed[action]:
        raise Rejected('Unknown or invalid options for action')
    metadata = {'format': suffix.lstrip('.'), 'size': len(data)}
    result = {'ok': True, 'status': 'ready', 'error': None, 'metadata': metadata,
              'truncated': False, 'truncation': [], 'warnings': []}

    def boolean(name, default):
        value = options.get(name, default)
        if type(value) is not bool:
            raise Rejected(name + ' must be boolean')
        return value

    def output(payload):
        if len(payload) > limits['expanded']:
            raise Rejected('Output size limit exceeded')
        result['output_data'] = base64.b64encode(payload).decode('ascii')
        metadata['output_format'] = output_suffix.lstrip('.')
        return result

    if action in ('stats', 'filter', 'sort', 'format', 'validate'):
        expected = ('.json',) if action in ('format', 'validate') else ('.csv', '.tsv')
        if suffix not in expected:
            return dict(result, ok=False, status='unsupported', error='Action requires ' + '/'.join(expected))
        encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        text = data.decode(encoding)
        if len(text) > limits['characters']:
            raise Rejected('Character limit exceeded; operation requires complete input')
        if '\x00' in text:
            raise Rejected('Binary content is not supported as text')
        if action in ('filter', 'sort', 'format') and output_suffix != suffix:
            raise Rejected('Destination extension must match input format')
        if suffix == '.json':
            obj = strict_json(text)
            metadata.update(valid=True, json_type=type(obj).__name__)
            if action == 'validate':
                return result
            indent = options.get('indent', 2)
            if type(indent) is not int or not 0 <= indent <= 8:
                raise Rejected('indent must be an integer from 0 to 8')
            rendered = json.dumps(obj, ensure_ascii=False, allow_nan=False, indent=indent,
                                  sort_keys=boolean('sort_keys', False)) + '\n'
            if len(rendered) > limits['characters']:
                raise Rejected('Formatted output character limit exceeded')
            return output(rendered.encode('utf-8'))
        delimiter = '\t' if suffix == '.tsv' else ','
        rows = []
        for row in csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True):
            if len(rows) >= limits['rows'] or len(row) > limits['columns']:
                raise Rejected('Table row or column limit exceeded')
            rows.append(row)
        if not rows or not rows[0] or any(not h.strip() for h in rows[0]) or len(set(rows[0])) != len(rows[0]):
            raise Rejected('Table requires unique nonempty column headers')
        headers, rows = rows[0], rows[1:]
        if any(len(row) != len(headers) for row in rows):
            raise Rejected('Table rows must match header width')
        metadata.update(rows=len(rows), columns=len(headers), headers=headers)

        def number(value):
            try:
                parsed = float(value)
            except (ValueError, TypeError, OverflowError):
                raise Rejected('Numeric operation requires finite numeric values') from None
            if not math.isfinite(parsed):
                raise Rejected('Numeric operation requires finite numeric values')
            return parsed

        if action == 'stats':
            statistics = []
            for index, name in enumerate(headers):
                values, missing, invalid = [], 0, 0
                for row in rows:
                    if not row[index].strip():
                        missing += 1
                        continue
                    try:
                        values.append(number(row[index]))
                    except Rejected:
                        invalid += 1
                try:
                    total = math.fsum(values)
                except OverflowError:
                    raise Rejected('Numeric aggregate exceeds supported range') from None
                statistics.append(dict(column=name, count=len(rows), numeric_count=len(values),
                                       missing=missing, nonnumeric=invalid,
                                       min=min(values) if values else None,
                                       max=max(values) if values else None,
                                       sum=total if values else None,
                                       mean=total / len(values) if values else None))
            result['statistics'] = statistics
            return result
        column = options.get('column')
        if not isinstance(column, str) or column not in headers:
            raise Rejected('column must name an existing header exactly')
        index = headers.index(column)
        if action == 'filter':
            condition = options.get('condition', 'equals')
            value = options.get('value')
            if condition not in ('equals', 'contains', 'gt', 'lt') or type(value) not in (str, int, float):
                raise Rejected('filter requires value and condition equals/contains/gt/lt')
            if isinstance(value, float) and not math.isfinite(value):
                raise Rejected('Filter value must be finite')
            if condition in ('gt', 'lt'):
                threshold = number(value)
                numeric = [(row, number(row[index])) for row in rows]
                rows = [row for row, n in numeric if (n > threshold if condition == 'gt' else n < threshold)]
            else:
                rows = [row for row in rows if (row[index] == str(value) if condition == 'equals' else str(value) in row[index])]
        else:
            numeric = boolean('numeric', False)
            rows = sorted(rows, key=lambda row: number(row[index]) if numeric else row[index],
                          reverse=not boolean('ascending', True))
        stream = io.StringIO(newline='')
        writer = csv.writer(stream, delimiter=delimiter, lineterminator='\n')
        writer.writerow(headers)
        writer.writerows(rows)
        rendered = stream.getvalue()
        if len(rendered) > limits['characters']:
            raise Rejected('Output character limit exceeded')
        metadata['output_rows'] = len(rows)
        result['warnings'] = ['Cell values are preserved as data; spreadsheet applications may interpret formula-like cells.']
        return output(rendered.encode('utf-8'))

    if suffix not in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp', '.ico'):
        return dict(result, ok=False, status='unsupported', error='Action requires a supported raster image')
    formats = {'.png': 'PNG', '.jpg': 'JPEG', '.jpeg': 'JPEG', '.webp': 'WEBP'}
    if output_suffix not in formats:
        raise Rejected('Image destination must be .png, .jpg, .jpeg or .webp')
    quality = options.get('quality', 90)
    if type(quality) is not int or not 1 <= quality <= 100:
        raise Rejected('quality must be an integer from 1 to 100')
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(data)) as source:
        if source.width * source.height > limits['pixels']:
            raise Rejected('Image pixel limit exceeded')
        if getattr(source, 'n_frames', 1) != 1:
            raise Rejected('Animated or multipage image transformations are unsupported')
        im = ImageOps.exif_transpose(source)
        metadata.update(width=im.width, height=im.height)
        if action == 'resize':
            width, height = options.get('width'), options.get('height')
            if width is None and height is None:
                raise Rejected('resize requires width or height')
            for dimension in (width, height):
                if dimension is not None and (type(dimension) is not int or not 1 <= dimension <= limits['pixels']):
                    raise Rejected('Dimensions must be bounded positive integers')
            width = width or max(1, round(im.width * height / im.height))
            height = height or max(1, round(im.height * width / im.width))
            if width * height > limits['pixels']:
                raise Rejected('Output image pixel limit exceeded')
            im = im.resize((width, height), Image.Resampling.LANCZOS)
        # Copy pixels into a clean image so source EXIF/ICC/comments are not forwarded.
        clean = Image.new('RGBA', im.size)
        clean.paste(im.convert('RGBA'))
        if formats[output_suffix] == 'JPEG':
            flattened = Image.new('RGB', clean.size, 'white')
            flattened.paste(clean, mask=clean.getchannel('A'))
            clean = flattened
        stream = io.BytesIO()
        clean.save(stream, format=formats[output_suffix], quality=quality)
        metadata.update(output_width=clean.width, output_height=clean.height,
                        metadata_removed=True)
        return output(stream.getvalue())


def archive(data, limits):
    z = zipfile.ZipFile(io.BytesIO(data))
    infos = z.infolist()
    if len(infos) > limits['members']:
        raise Rejected('Archive member limit exceeded')
    total, seen, entries = 0, set(), []
    for item in infos:
        name = item.filename
        parts = name.rstrip('/').split('/')
        mode = item.external_attr >> 16
        if (not name or '\\' in name or '\x00' in name or name.startswith('/')
                or any(p in ('', '.', '..') or ':' in p or p.endswith((' ', '.'))
                       or re.match(r'(?i)^(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\.|$)', p)
                       for p in parts)
                or stat.S_IFMT(mode) not in (0, stat.S_IFREG, stat.S_IFDIR)):
            raise Rejected('Unsafe archive member')
        key = '/'.join(parts).casefold()
        if key in seen:
            raise Rejected('Duplicate archive member')
        seen.add(key)
        if item.flag_bits & 1:
            raise Rejected('Encrypted archive is unsupported')
        total += item.file_size
        if total > limits['expanded']:
            raise Rejected('Archive expanded size limit exceeded')
        entries.append((item, '/'.join(parts)))
    files = {name.casefold() for item, name in entries if not item.is_dir()}
    for item, name in entries:
        if any(str(p).casefold() in files for p in PurePosixPath(name).parents if str(p) != '.'):
            raise Rejected('Archive file/directory collision')
    return z, entries


def parse(data, suffix, action, limits, options=None, output_suffix=None):
    if action == '_analysis_image':
        if len(data)>limits['input_bytes']:raise Rejected('Input size limit exceeded')
        from PIL import Image, ImageOps
        with Image.open(io.BytesIO(data)) as source:
            if source.width*source.height>limits['pixels']:raise Rejected('Image pixel limit exceeded')
            image=ImageOps.exif_transpose(source).convert('RGB')
            image.thumbnail((1600,1600))
            # Re-encode pixels only; do not transmit EXIF or other file metadata.
            output=io.BytesIO();image.save(output,format='JPEG',quality=85)
            return {'ok':True,'status':'ready','metadata':{'width':image.width,'height':image.height},
                    'image_data':base64.b64encode(output.getvalue()).decode('ascii')}
    if len(data) > limits['input_bytes']:
        raise Rejected('Input size limit exceeded')
    suffix = suffix.lower()
    if action in ('stats', 'filter', 'sort', 'format', 'validate', 'resize', 'convert'):
        return transform(data, suffix, action, limits, options, output_suffix)
    metadata = {'format': suffix.lstrip('.'), 'size': len(data)}
    reasons, chunks, length = [], [], 0
    sections = []

    def section(kind, label):
        sections.append({'kind': kind, 'label': str(label), 'start': length})

    def add(value):
        nonlocal length
        value = str(value)
        remaining = limits['characters'] - length
        if len(value) > remaining and 'characters' not in reasons:
            reasons.append('characters')
        if remaining:
            chunks.append(value[:remaining])
        length += min(len(value), remaining)

    if suffix == '.zip' or suffix in ('.docx', '.pptx', '.xlsx'):
        z, entries = archive(data, limits)
        z.close()
    if suffix == '.zip':
        metadata['members'] = len(entries)
        result = []
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for item, name in entries:
                payload = b'' if item.is_dir() or action == 'inspect' else z.read(item)
                result.append({'name': name, 'directory': item.is_dir(), 'data': base64.b64encode(payload).decode()})
        return {'ok': True, 'status': 'ready', 'error': None, 'metadata': metadata, 'members': result}
    if suffix in TEXT_FORMATS:
        encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        text = data.decode(encoding)
        if '\x00' in text:
            raise Rejected('Binary content is not supported as text')
        metadata['encoding'] = encoding
        if suffix == '.json':
            obj = strict_json(text)
            metadata['json_type'] = type(obj).__name__
            metadata['valid'] = True
            if isinstance(obj, (dict, list)):
                metadata['items'] = len(obj)
        if suffix == '.xml':
            if re.search(r'<!\s*(DOCTYPE|ENTITY)\b', text, re.I):
                raise Rejected('XML DTDs and entities are unsupported')
            root = ElementTree.fromstring(text)
            metadata.update(valid=True, root=root.tag)
        if suffix in ('.csv', '.tsv'):
            rows = csv.reader(io.StringIO(text), delimiter='\t' if suffix == '.tsv' else ',')
            count = 0
            columns = 0
            for row in rows:
                if count >= limits['rows']:
                    reasons.append('rows')
                    break
                add('\t'.join(row[:limits['columns']]) + '\n')
                if len(row) > limits['columns'] and 'columns' not in reasons:
                    reasons.append('columns')
                count += 1
                columns = max(columns, len(row))
            metadata['rows_read'] = count
            metadata['columns_read'] = min(columns, limits['columns'])
        else:
            add(text)
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise Rejected('Encrypted PDF is unsupported')
        metadata['pages'] = len(reader.pages)
        if len(reader.pages) > limits['pages']:
            reasons.append('pages')
        for index, page in enumerate(reader.pages[:limits['pages']], 1):
            section('page', index)
            add((page.extract_text() or '') + '\n')
        metadata['pages_read'] = len(sections)
    elif suffix == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        metadata.update(paragraphs=len(document.paragraphs), tables=len(document.tables))
        for p in document.paragraphs[:limits['rows']]:
            add(p.text + '\n')
        if len(document.paragraphs) > limits['rows']:
            reasons.append('rows')
        count = 0
        for table in document.tables:
            for row in table.rows:
                count += 1
                if count > limits['rows']:
                    break
                add('\t'.join(c.text for c in row.cells[:limits['columns']]) + '\n')
                if len(row.cells) > limits['columns'] and 'columns' not in reasons:
                    reasons.append('columns')
            if count > limits['rows']:
                reasons.append('rows')
                break
    elif suffix == '.pptx':
        from pptx import Presentation
        slides = Presentation(io.BytesIO(data)).slides
        metadata['slides'] = len(slides)
        if len(slides) > limits['slides']:
            reasons.append('slides')
        for i, slide in enumerate(slides):
            if i >= limits['slides']:
                break
            section('slide', i + 1)
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text + '\n')
                if shape.has_table:
                    for row in list(shape.table.rows)[:limits['rows']]:
                        add('\t'.join(c.text for c in list(row.cells)[:limits['columns']]) + '\n')
                        if len(row.cells) > limits['columns']:
                            reasons.append('columns')
                    if len(shape.table.rows) > limits['rows']:
                        reasons.append('rows')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            metadata['sheets'] = len(book.sheetnames)
            metadata['sheet_names'] = book.sheetnames[:limits['sheets']]
            if len(book.sheetnames) > limits['sheets']:
                reasons.append('sheets')
            for sheet in book.worksheets[:limits['sheets']]:
                section('sheet', sheet.title)
                add(sheet.title + '\n')
                if (sheet.max_row or 0) > limits['rows']:
                    reasons.append('rows')
                if (sheet.max_column or 0) > limits['columns']:
                    reasons.append('columns')
                for row in sheet.iter_rows(max_row=min(sheet.max_row or limits['rows'], limits['rows']), max_col=min(sheet.max_column or 1, limits['columns']), values_only=True):
                    add('\t'.join('' if c is None else str(c) for c in row) + '\n')
        finally:
            book.close()
    elif suffix in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff', '.webp', '.ico'):
        from PIL import Image
        with Image.open(io.BytesIO(data)) as im:
            if im.width * im.height > limits['pixels']:
                raise Rejected('Image pixel limit exceeded')
            metadata.update(width=im.width, height=im.height, mode=im.mode, image_format=im.format)
            im.verify()
    else:
        return {'ok': False, 'status': 'unsupported', 'error': 'Unsupported file format',
                'metadata': metadata,
                'hint': 'Use DOCX, PPTX or XLSX for legacy Office files; media conversion and OCR are not configured.'}
    text = ''.join(chunks)
    metadata.update(characters_read=len(text), words_read=len(text.split()),
                    lines_read=len(text.splitlines()), counts_complete=not reasons)
    for index, item in enumerate(sections):
        item['end'] = sections[index + 1]['start'] if index + 1 < len(sections) else length
    return {'ok': True, 'status': 'ready', 'error': None, 'metadata': metadata,
            'text': '' if action == 'inspect' else text, 'sections': sections,
            'warnings': ['No extractable text; OCR is not configured.'] if suffix == '.pdf' and not text.strip() else [],
            'truncated': bool(reasons), 'truncation': sorted(set(reasons))}


def main():
    request = json.load(sys.stdin)
    try:
        result = parse(base64.b64decode(request['data']), request['suffix'], request['action'], request['limits'],
                       request.get('options'), request.get('output_suffix'))
    except ImportError as exc:
        result = {'ok': False, 'status': 'unavailable', 'error': 'Optional parser dependency missing', 'dependency': exc.name}
    except Rejected as exc:
        result = {'ok': False, 'status': 'rejected', 'error': str(exc)}
    except UnicodeError:
        result = {'ok': False, 'status': 'failed', 'error': 'Unsupported text encoding; use UTF-8 or BOM-marked UTF-16', 'code': 'TEXT_ENCODING'}
    except Exception:
        result = {'ok': False, 'status': 'failed', 'error': 'Malformed or unreadable document'}
    json.dump(result, sys.stdout, ensure_ascii=True)


if __name__ == '__main__':
    main()
